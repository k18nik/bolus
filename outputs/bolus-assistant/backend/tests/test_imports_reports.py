import io,json,zipfile
from datetime import datetime,timezone,timedelta
from zoneinfo import ZoneInfo
from pathlib import Path
import pytest
from openpyxl import load_workbook
from app.imports.apple_health import parse_apple_health
from app.reports.generator import safe_cell

XML=b'''<?xml version="1.0" encoding="UTF-8"?><HealthData><Workout workoutActivityType="HKWorkoutActivityTypeWalking" startDate="2026-09-30 15:00:00 +0300" endDate="2026-09-30 15:32:00 +0300" sourceName="Apple Watch"/><ActivitySummary dateComponents="2026-09-30" activeEnergyBurned="250" activeEnergyBurnedUnit="kcal" appleExerciseTime="32"/></HealthData>'''

def test_health_parser_zip_and_entities():
    rows=parse_apple_health(XML);assert len(rows)==2
    assert rows[0]['data']['duration_minutes']==32
    assert rows[0]['occurred_at']=='2026-09-30T12:00:00+00:00'
    buf=io.BytesIO()
    with zipfile.ZipFile(buf,'w') as z:z.writestr('apple_health_export/export.xml',XML)
    assert parse_apple_health(buf.getvalue())==rows
    with pytest.raises(ValueError):parse_apple_health(b'<!DOCTYPE x [<!ENTITY a "malicious">]><HealthData/>')
    with pytest.raises(ValueError):parse_apple_health(b'<HealthData/>')

def test_health_import_dedup(demo):
    a=demo.post('/api/imports/apple-health',files={'file':('export.xml',XML,'text/xml')})
    assert a.status_code==200,a.text
    assert a.json()['inserted']==2
    b=demo.post('/api/imports/apple-health',files={'file':('export.xml',XML,'text/xml')})
    assert b.json()['inserted']==0 and b.json()['skipped']==2
    rows=demo.get('/api/diary?date_from=2026-09-30&date_to=2026-09-30').json()
    assert any(e['kind']=='activity' and e['data'].get('source')=='apple_health' for e in rows)

@pytest.mark.parametrize('format',['pdf','xlsx','csv','json'])
def test_report_formats(demo,format):
    today=datetime.now(ZoneInfo('Europe/Moscow')).date()
    response=demo.post('/api/reports',json={'type':'doctor','format':format,'date_from':str(today-timedelta(days=6)),'date_to':str(today)})
    assert response.status_code==202,response.text
    job_id=response.json()['id'];jobs=demo.get('/api/reports').json();job=next(j for j in jobs if j['id']==job_id)
    assert job['status']=='completed',job
    content=demo.get('/api/reports/'+job_id+'/download').content
    if format=='pdf':assert content.startswith(b'%PDF') and len(content)>5000
    if format=='xlsx':
        wb=load_workbook(io.BytesIO(content));assert {'Summary','Glucose','Insulin','Meals','Meal Items','Activities','Cycle','Bolus Calculations','Therapy Profiles','Patterns'}<=set(wb.sheetnames)
        summary=dict(zip(next(wb['Summary'].values),list(wb['Summary'].values)[1]));analytics=demo.get('/api/analytics?days=7').json()['metrics']
        assert summary['mean_glucose']==analytics['mean_glucose']
        assert '+03:00' in wb['Glucose'].cell(2,2).value
    if format=='csv':
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            text=z.read('glucose.csv').decode('utf-8-sig');assert 'value_mmol' in text and '+03:00' in text
    if format=='json':
        data=json.loads(content);assert data['schema_version']=='1.0' and data['exported_at'] and data['entries']
        assert 'password_hash' not in content.decode()
    assert demo.delete('/api/reports/'+job_id).status_code==200
    assert demo.get('/api/reports/'+job_id+'/download').status_code==404

def test_csv_formula_injection():
    assert safe_cell('=cmd()')=="'=cmd()"
    assert safe_cell('@SUM(1,2)')=="'@SUM(1,2)"
    assert safe_cell(2)==2

def test_report_invalid_period(demo):
    assert demo.post('/api/reports',json={'type':'doctor','format':'pdf','date_from':'2026-10-02','date_to':'2026-01-01'}).status_code==422

def test_report_mgdl_period_options(demo):
    user=demo.get('/api/users/me').json()
    settings={k:user[k] for k in ('name','timezone','glucose_unit','theme_id','mascot_id')};settings['glucose_unit']='mg/dL'
    assert demo.patch('/api/users/me',json=settings).status_code==200
    today=datetime.now(ZoneInfo('Europe/Moscow')).date()
    job=demo.post('/api/reports',json={'type':'doctor','format':'xlsx','date_from':str(today-timedelta(days=6)),'date_to':str(today),'include_cycle':False,'include_nutrition':False}).json()
    wb=load_workbook(io.BytesIO(demo.get('/api/reports/'+job['id']+'/download').content))
    assert 'Cycle' not in wb.sheetnames and 'Meals' not in wb.sheetnames
    summary=dict(zip(next(wb['Summary'].values),list(wb['Summary'].values)[1]))
    avg=demo.get('/api/analytics?days=7').json()['metrics']['mean_glucose']
    assert summary['glucose_unit']=='mg/dL' and summary['mean_glucose']==round(avg*18,2)
