import csv, io, json, os, zipfile
from pathlib import Path
from datetime import datetime,timezone,timedelta,date
from zoneinfo import ZoneInfo
from xml.sax.saxutils import escape
from sqlalchemy import select
from app.config import settings
from app.db import SessionLocal
from app.models.entities import User,Entry,Profile,Calculation,Cycle,CustomFood,ReportJob,AuditLog,FoodFavorite,AIInsight,now
from app.repositories.diary import period_entries,entry_dict
from app.analytics.engine import summarize,hourly_profile

SHEETS={'Glucose':'glucose','Insulin':'insulin','Meals':'meal','Activities':'activity','Activity Summaries':'activity_summary'}
def safe_cell(value):
    if isinstance(value,(dict,list)): value=json.dumps(value,ensure_ascii=False)
    if isinstance(value,str) and value.startswith(('=','+','-','@','\t','\r')): return "'"+value
    return value

def tables_for(entries,profiles,calculations,cycles,foods,tz,unit):
    tables={}
    for name,kind in SHEETS.items():
        records=[]
        for e in entries:
            if e['kind']!=kind: continue
            data=dict(e['data']);data.pop('items',None)
            if kind=='glucose':
                data['value']=round(data['value_mmol']*(18 if unit=='mg/dL' else 1),2);data['unit']=unit
            records.append({'id':e['id'],'timestamp':datetime.fromisoformat(e['occurred_at']).astimezone(ZoneInfo(tz)).isoformat(),**data})
        tables[name]=records
    tables['Meal Items']=[{'meal_id':e['id'],**i} for e in entries if e['kind']=='meal' for i in e['data'].get('items',[])]
    tables['Cycle']=[{'id':c.id,'start_date':c.start_date,'end_date':c.end_date,'cycle_length':c.cycle_length} for c in cycles]
    tables['Bolus Calculations']=[{'id':c.id,'calculated_at':c.calculated_at,'algorithm_version':c.algorithm_version,'actual_bolus':c.actual_bolus,'input_snapshot':c.input_snapshot,'calculation_snapshot':c.calculation_snapshot} for c in calculations]
    tables['Therapy Profiles']=[{'id':p.id,'version':p.version,'valid_from':p.valid_from,'valid_to':p.valid_to,'source':p.source,'data':p.data} for p in profiles]
    tables['Foods']=[{'id':f.id,'name':f.name,'is_recipe':f.is_recipe,'data':f.data} for f in foods]
    tables['Notes']=[{'id':e['id'],'timestamp':e['occurred_at'],**e['data']} for e in entries if e['kind']=='note']
    tables['Patterns']=[]
    return tables

def columns(rows): return list(dict.fromkeys(k for row in rows for k in row)) or ['No data']
def csv_bytes(rows):
    out=io.StringIO(); writer=csv.DictWriter(out,fieldnames=columns(rows));writer.writeheader()
    for row in rows: writer.writerow({k:safe_cell(v) for k,v in row.items()})
    return out.getvalue().encode('utf-8-sig')


def generate_report(job_id):
    with SessionLocal() as db:
        job=db.get(ReportJob,job_id)
        if not job or job.status=='completed': return
        user=db.get(User,job.user_id)
        if not user:return
        owner_id=user.id; job.status='processing';job.progress=10;db.commit()
        try:
            df=date.fromisoformat(job.date_from);dt=date.fromisoformat(job.date_to)
            entries=period_entries(db,user.id,df,dt,user.timezone)
            metrics=summarize(entries,(dt-df).days+1,user.timezone)
            profiles=list(db.scalars(select(Profile).where(Profile.user_id==user.id).order_by(Profile.version)))
            calculations=list(db.scalars(select(Calculation).where(Calculation.user_id==user.id)))
            calculations=[c for c in calculations if df<=datetime.fromisoformat(c.calculated_at).astimezone(ZoneInfo(user.timezone)).date()<=dt]
            cycles=list(db.scalars(select(Cycle).where(Cycle.user_id==user.id)))
            foods=list(db.scalars(select(CustomFood).where(CustomFood.user_id==user.id)))
            cycles_in_period=[c for c in cycles if c.start_date<=str(dt) and date.fromisoformat(c.start_date)+timedelta(days=c.cycle_length-1)>=df]
            tables=tables_for(entries,profiles,calculations,cycles_in_period,foods,user.timezone,user.glucose_unit)
            if not job.options.get('include_nutrition',True):
                for name in ('Meals','Meal Items','Foods'):tables.pop(name,None)
            if not job.options.get('include_cycle',True):tables.pop('Cycle',None)
            out=Path(settings.report_dir).resolve();out.mkdir(parents=True,exist_ok=True)
            extension='zip' if job.format=='csv' else job.format
            path=out/f'{job.id}.{extension}'
            if job.format=='json':
                # Full backup always includes all rows, independent of selected report period.
                all_entries=[entry_dict(e) for e in db.scalars(select(Entry).where(Entry.user_id==user.id))]
                all_calcs=list(db.scalars(select(Calculation).where(Calculation.user_id==user.id)))
                backup=tables_for(all_entries,profiles,all_calcs,cycles,foods,user.timezone,user.glucose_unit)
                backup['Food Favorites']=[{'id':f.id,'food_key':f.food_key,'data':f.data} for f in db.scalars(select(FoodFavorite).where(FoodFavorite.user_id==user.id))]
                backup['AI Insights']=[{k:getattr(r,k) for k in ('id','question','response','context','provider','model','usage','calculation_id','created_at')} for r in db.scalars(select(AIInsight).where(AIInsight.user_id==user.id))]
                backup['Audit']=[{'action':a.action,'entity_type':a.entity_type,'entity_id':a.entity_id,'old_value':a.old_value,'new_value':a.new_value,'timestamp':a.timestamp} for a in db.scalars(select(AuditLog).where(AuditLog.user_id==user.id))]
                path.write_text(json.dumps({'schema_version':'1.0','exported_at':now(),'user':{k:getattr(user,k) for k in ('id','email','name','timezone','locale','glucose_unit','weight_unit','theme_id','mascot_id','created_at')},'entries':all_entries,'datasets':backup},ensure_ascii=False,indent=2),encoding='utf-8')
            elif job.format=='csv':
                with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as archive:
                    for name,rows in tables.items():archive.writestr(name.lower().replace(' ','_')+'.csv',csv_bytes(rows))
                    archive.writestr('summary.csv',csv_bytes([metrics]))
            elif job.format=='xlsx':
                from openpyxl import Workbook
                from openpyxl.styles import Font,PatternFill
                wb=Workbook();wb.remove(wb.active)
                for name,rows in {'Summary':[{'period':f'{df} — {dt}','timezone':user.timezone,'glucose_unit':user.glucose_unit,**{k:(round(v*18,2) if user.glucose_unit=='mg/dL' and k in ('mean_glucose','median_glucose','min','max','standard_deviation') and v is not None else v) for k,v in metrics.items()}}],**tables}.items():
                    ws=wb.create_sheet(name);keys=columns(rows);ws.append(keys)
                    for row in rows: ws.append([safe_cell(row.get(k)) for k in keys])
                    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
                    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='237F6C')
                    for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=24
                wb.save(path)
            else:
                from app.reports.pdf_renderer import render_pdf
                render_pdf(path,user,job,metrics,entries,profiles,cycles)
            # Account deletion or job deletion while processing must not resurrect orphan files.
            db.expire_all()
            if not db.get(User,owner_id) or not db.get(ReportJob,job_id):path.unlink(missing_ok=True);return
            job.file_path=str(path);job.status='completed';job.progress=100;job.completed_at=now();job.expires_at=(datetime.now(timezone.utc)+timedelta(days=7)).isoformat();db.commit()
        except Exception:
            if 'path' in locals():path.unlink(missing_ok=True)
            db.rollback();job=db.get(ReportJob,job_id)
            if job:job.status='failed';job.error='Не удалось создать отчёт. Попробуйте ещё раз.';db.commit()
            raise
